import tempfile
import unittest
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image as PillowImage

from ingest import SCHEMA, as_number, canonical_for, excel_products, import_file, infer_brand_from_filename, is_size_header, parse_classic_page, parse_destock_page, parse_lot_page, parse_offer_page, sync_excel_images, xlsx_embedded_images


class ImporterTests(unittest.TestCase):
    def test_excel_header_detection_and_mapping(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["Supplier offer"])
            sheet.append(["Brand", "Model", "SKU", "Qty", "Net Price", "Retail Price"])
            sheet.append(["Nike", "Pegasus", "ABC-123", 8, 25, 80])
            book.save(path)
            rows = list(excel_products(path))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].brand, "Nike")
        self.assertEqual(rows[0].sku, "ABC-123")
        self.assertEqual(rows[0].quantity, 8)
        self.assertEqual(rows[0].wholesale_price, 25)

    def test_offer_pdf_pattern(self):
        rows = parse_offer_page("J30J320935-BEH\nCORE TEE\nWSP £18.70\nRRP £39.90", 2)
        self.assertEqual(rows[0].sku, "J30J320935-BEH")
        self.assertEqual(rows[0].retail_price, 39.9)

    def test_destock_pdf_pattern(self):
        text = "RUNNER GW0005 - CHAUSSURE Ref : Gamme : 40 2 Total 12\nP.Tarif 13.00 €\nRRP 35.00 €"
        rows = parse_destock_page(text, 1)
        self.assertEqual(rows[0].sku, "GW0005")
        self.assertEqual(rows[0].quantity, 12)

    def test_take_all_lot_pattern(self):
        text = "ULTRA PLAY10768901 - CHAUSSUREEURRef :Gamme :391401Total12"
        rows = parse_lot_page(text, 1, 18)
        self.assertEqual(rows[0].sku, "10768901")
        self.assertEqual(rows[0].quantity, 12)
        self.assertEqual(rows[0].wholesale_price, 18)

    def test_brand_from_single_brand_filename(self):
        self.assertEqual(infer_brand_from_filename("New Offer Tommy Hilfiger.xlsx"), "TOMMY HILFIGER")
        self.assertIsNone(infer_brand_from_filename("Lot Shoes Women.pdf"))

    def test_supplier_price_header_edge_cases(self):
        wholesale = ("Price (EURO)", "your price euro", "WHS PRICE (ZI02)", "Net Prices € MOQ 100pcs", "Unit Price ex VAT")
        retail = ("SRP PRICE", "Public price", "RRP/UVP", "Listing Price")
        for header in wholesale:
            self.assertEqual(canonical_for(header), "wholesale_price", header)
        for header in retail:
            self.assertEqual(canonical_for(header), "retail_price", header)
        self.assertEqual(canonical_for("Total RRP"), "total_retail_value")
        self.assertIsNone(canonical_for("Amount (GBP)"))

    def test_force_reimport_preserves_product_id(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "prices.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["SKU", "Price (EURO)", "RRP"])
            sheet.append(["ABC-123", 20, 50])
            book.save(path)
            connection = sqlite3.connect(":memory:")
            connection.executescript(SCHEMA)
            import_file(connection, path, default_currency="EUR")
            first = connection.execute("SELECT id FROM products").fetchone()[0]
            import_file(connection, path, force=True, default_currency="EUR")
            second = connection.execute("SELECT id FROM products").fetchone()[0]
            connection.close()
        self.assertEqual(first, second)

    def test_ean_helper_sheet_enriches_variants_without_duplicates(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "matrix.xlsx"
            book = Workbook()
            main = book.active
            main.title = "Offer"
            main.append(["Article", "Model", "C38", "C40", "Total"])
            main.append(["ABC-1", "Work trousers", 3, 2, 5])
            ean = book.create_sheet("EAN HS Code")
            ean.append(["Article", "C38", "C40", "Country", "Material", "HS Code"])
            ean.append(["ABC-1", "1111111111111", "2222222222222", "Italy", "Cotton", "620342"])
            book.save(path)
            rows = list(excel_products(path))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].origin, "Italy")
        self.assertEqual(rows[0].hs_code, "620342")
        self.assertEqual(rows[0].material, "Cotton")
        self.assertEqual(len(rows[0].variants), 2)
        self.assertEqual(rows[0].variants[0]["barcode"], "1111111111111")

    def test_stock_threshold_and_alpha_size(self):
        self.assertEqual(as_number("60+"), 60)
        self.assertTrue(is_size_header("XXL"))

    def test_excel_image_anchor_is_associated_with_product_row(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            png = root / "photo.png"
            PillowImage.new("RGB", (12, 12), "red").save(png)
            workbook_path = root / "with-image.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Stock"
            sheet.append(["Image", "SKU", "Title"])
            sheet.append([None, "ABC-123", "Red product"])
            sheet.add_image(ExcelImage(png), "A2")
            book.save(workbook_path)
            anchors = list(xlsx_embedded_images(workbook_path))
            self.assertEqual((anchors[0]["sheet"], anchors[0]["row"]), ("Stock", 2))
            connection = sqlite3.connect(":memory:")
            connection.executescript(SCHEMA)
            import_file(connection, workbook_path, default_currency="EUR")
            source_id = connection.execute("SELECT id FROM source_files").fetchone()[0]
            matched, unmatched = sync_excel_images(connection, source_id, workbook_path, root / "media", root)
            image_row = connection.execute("SELECT source_row,match_method FROM product_images").fetchone()
            connection.close()
        self.assertEqual((matched, unmatched), (1, 0))
        self.assertEqual(image_row, (2, "excel_cell_anchor"))

    def test_classic_pdf_pattern(self):
        text = "ARIZONA\nNATURAL LEATHER\nBLACK\n1026693\nREGULAR: 1026693\nNARROW: 1026719\n140.00 35 - 43"
        rows = parse_classic_page(text, 1)
        self.assertEqual(rows[0].sku, "1026693")
        self.assertEqual(rows[0].size, "35 - 43")


if __name__ == "__main__":
    unittest.main()
