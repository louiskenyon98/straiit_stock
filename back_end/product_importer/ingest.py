#!/usr/bin/env python3
"""Import product spreadsheets and PDF catalogues into a SQLite database.

The importer deliberately keeps both a normalized set of website-friendly fields
and every original supplier field in ``attributes_json``.  This makes it useful
with mixed supplier formats without throwing information away.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import sqlite3
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Iterator
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from pypdf import PdfReader


SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS source_files (
    id INTEGER PRIMARY KEY,
    path TEXT NOT NULL UNIQUE,
    filename TEXT NOT NULL,
    file_type TEXT NOT NULL,
    sha256 TEXT NOT NULL,
    size_bytes INTEGER NOT NULL,
    modified_at TEXT NOT NULL,
    imported_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    product_count INTEGER NOT NULL DEFAULT 0,
    warning TEXT
);
CREATE TABLE IF NOT EXISTS products (
    id INTEGER PRIMARY KEY,
    source_file_id INTEGER NOT NULL REFERENCES source_files(id) ON DELETE CASCADE,
    source_location TEXT NOT NULL,
    source_key TEXT NOT NULL,
    brand TEXT,
    name TEXT,
    category TEXT,
    model TEXT,
    sku TEXT,
    barcode TEXT,
    description TEXT,
    color TEXT,
    size TEXT,
    gender TEXT,
    material TEXT,
    quantity REAL,
    wholesale_price REAL,
    retail_price REAL,
    currency TEXT,
    status TEXT,
    image_path TEXT,
    offer_type TEXT,
    origin TEXT,
    hs_code TEXT,
    grade TEXT,
    model_year TEXT,
    collection TEXT,
    frame_color TEXT,
    lens_color TEXT,
    color_code TEXT,
    material_code TEXT,
    frame_material TEXT,
    bridge_size TEXT,
    branch_size TEXT,
    uva_filter TEXT,
    fitting TEXT,
    release_code TEXT,
    phase TEXT,
    order_quantity REAL,
    exceeding_quantity REAL,
    total_retail_value REAL,
    attributes_json TEXT NOT NULL DEFAULT '{}',
    raw_text TEXT,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    UNIQUE(source_file_id, source_key)
);
CREATE INDEX IF NOT EXISTS idx_products_brand ON products(brand);
CREATE INDEX IF NOT EXISTS idx_products_category ON products(category);
CREATE INDEX IF NOT EXISTS idx_products_sku ON products(sku);
CREATE INDEX IF NOT EXISTS idx_products_barcode ON products(barcode);
CREATE INDEX IF NOT EXISTS idx_products_name ON products(name);
CREATE TABLE IF NOT EXISTS product_variants (
    id INTEGER PRIMARY KEY,
    product_id INTEGER NOT NULL REFERENCES products(id) ON DELETE CASCADE,
    size TEXT NOT NULL,
    quantity REAL,
    barcode TEXT,
    sku TEXT,
    attributes_json TEXT NOT NULL DEFAULT '{}',
    UNIQUE(product_id, size)
);
CREATE INDEX IF NOT EXISTS idx_variants_product ON product_variants(product_id);
CREATE INDEX IF NOT EXISTS idx_variants_barcode ON product_variants(barcode);
CREATE TABLE IF NOT EXISTS product_images (
    id INTEGER PRIMARY KEY,
    product_id INTEGER NOT NULL REFERENCES products(id) ON DELETE CASCADE,
    source_file_id INTEGER NOT NULL REFERENCES source_files(id) ON DELETE CASCADE,
    image_path TEXT NOT NULL,
    mime_type TEXT,
    source_sheet TEXT,
    source_row INTEGER,
    source_page INTEGER,
    source_x REAL,
    source_y REAL,
    width REAL,
    height REAL,
    sha256 TEXT NOT NULL,
    sort_order INTEGER NOT NULL DEFAULT 0,
    is_primary INTEGER NOT NULL DEFAULT 0,
    match_method TEXT,
    UNIQUE(product_id, image_path)
);
CREATE INDEX IF NOT EXISTS idx_product_images_product ON product_images(product_id);
CREATE INDEX IF NOT EXISTS idx_product_images_source ON product_images(source_file_id);
CREATE INDEX IF NOT EXISTS idx_product_images_hash ON product_images(sha256);
CREATE VIEW IF NOT EXISTS website_products AS
SELECT p.id, p.brand, p.name, p.category, p.model, p.sku, p.barcode,
       p.description, p.color, p.size, p.gender, p.material, p.quantity,
       p.wholesale_price, p.retail_price, p.currency, p.status, p.image_path,
       p.offer_type, p.origin, p.hs_code, p.grade, p.model_year, p.collection,
       p.frame_color, p.lens_color, p.color_code, p.material_code,
       p.frame_material, p.bridge_size, p.branch_size, p.uva_filter, p.fitting,
       p.release_code, p.phase, p.order_quantity, p.exceeding_quantity,
       p.total_retail_value,
       s.filename AS source_file, p.source_location, p.attributes_json
FROM products p JOIN source_files s ON s.id = p.source_file_id;
"""

PRODUCT_COLUMN_MIGRATIONS = {
    "offer_type": "TEXT", "origin": "TEXT", "hs_code": "TEXT", "grade": "TEXT",
    "model_year": "TEXT", "collection": "TEXT", "frame_color": "TEXT", "lens_color": "TEXT",
    "color_code": "TEXT", "material_code": "TEXT", "frame_material": "TEXT",
    "bridge_size": "TEXT", "branch_size": "TEXT", "uva_filter": "TEXT", "fitting": "TEXT",
    "release_code": "TEXT", "phase": "TEXT", "order_quantity": "REAL",
    "exceeding_quantity": "REAL", "total_retail_value": "REAL",
}


def ensure_schema(connection: sqlite3.Connection) -> None:
    """Apply additive migrations and refresh the website view for existing DBs."""
    existing = {row[1] for row in connection.execute("PRAGMA table_info(products)")}
    for column, sql_type in PRODUCT_COLUMN_MIGRATIONS.items():
        if column not in existing:
            connection.execute(f"ALTER TABLE products ADD COLUMN {column} {sql_type}")
    connection.execute("DROP VIEW IF EXISTS website_products")
    connection.execute("""
        CREATE VIEW website_products AS
        SELECT p.id, p.brand, p.name, p.category, p.model, p.sku, p.barcode,
               p.description, p.color, p.size, p.gender, p.material, p.quantity,
               p.wholesale_price, p.retail_price, p.currency, p.status, p.image_path,
               p.offer_type, p.origin, p.hs_code, p.grade, p.model_year, p.collection,
               p.frame_color, p.lens_color, p.color_code, p.material_code,
               p.frame_material, p.bridge_size, p.branch_size, p.uva_filter, p.fitting,
               p.release_code, p.phase, p.order_quantity, p.exceeding_quantity,
               p.total_retail_value, s.filename AS source_file, p.source_location,
               p.attributes_json
        FROM products p JOIN source_files s ON s.id = p.source_file_id
    """)
    connection.commit()


ALIASES = {
    "brand": ("brand", "marque", "manufacturer"),
    "name": ("title", "product", "product name", "item name", "designation", "style name"),
    "category": ("category", "type", "family", "product type", "gamme", "sun opth"),
    "model": ("model", "model code", "model color size", "model ref", "maison model name", "style", "model and size", "model and sizes"),
    "sku": ("sku", "item code", "product code", "reference", "ref", "article", "artilce", "style code", "size and col code"),
    "barcode": ("barcode", "upc", "upc code", "upc sku", "ean", "ean code", "gtin"),
    "description": ("description", "color description", "color descr", "details"),
    "color": ("color", "colour"),
    "size": ("size", "size range"),
    "gender": ("gender", "sex", "department", "target"),
    "material": ("material", "composition"),
    "quantity": ("quantity", "qty", "q ty", "qty avl", "stock", "available", "available bales", "total quantity", "total", "total pairs"),
    "wholesale_price": ("your price", "net price", "wholesale price", "wsp", "whs", "price per pair eur", "price"),
    "retail_price": ("retail price", "listing price", "public price", "rrp", "msrp"),
    "status": ("status", "item status", "availability"),
    "offer_type": ("offer", "offer type"),
    "origin": ("origin", "country", "country of origin", "made in"),
    "hs_code": ("hs code", "commodity code", "tariff code"),
    "grade": ("grade",),
    "model_year": ("year", "model year", "release year"),
    "collection": ("collection",),
    "frame_color": ("color frame", "frame color", "frame colour"),
    "lens_color": ("color lens", "lens color", "lens colour"),
    "color_code": ("col code", "color code", "colour code", "maison color code", "lens code"),
    "material_code": ("material code",),
    "frame_material": ("frame material",),
    "bridge_size": ("bridge size",),
    "branch_size": ("branch size", "temple length"),
    "uva_filter": ("uva filters", "uva filter"),
    "fitting": ("fitting",),
    "release_code": ("release code",),
    "phase": ("phase",),
    "order_quantity": ("order", "order qty", "ordered quantity"),
    "exceeding_quantity": ("exceeding qty", "excess quantity"),
    "total_retail_value": ("rrp total", "total rrp", "retail total", "total retail"),
}

SKIP_SHEETS = {"condition", "conditions", "terms", "notes", "read me", "instructions"}
PRODUCT_CODE = re.compile(r"\b(?=[A-Z0-9./_-]{5,}\b)(?=[A-Z0-9./_-]*[A-Z])(?=[A-Z0-9./_-]*\d)[A-Z0-9]+(?:[./_-]?[A-Z0-9]+)+\b", re.I)
MONEY = r"(?:£|€|\$|\ufffd)?\s*(\d+(?:[.,]\d{1,2})?)"


@dataclass
class Product:
    source_location: str
    source_key: str
    brand: str | None = None
    name: str | None = None
    category: str | None = None
    model: str | None = None
    sku: str | None = None
    barcode: str | None = None
    description: str | None = None
    color: str | None = None
    size: str | None = None
    gender: str | None = None
    material: str | None = None
    quantity: float | None = None
    wholesale_price: float | None = None
    retail_price: float | None = None
    currency: str | None = None
    status: str | None = None
    image_path: str | None = None
    offer_type: str | None = None
    origin: str | None = None
    hs_code: str | None = None
    grade: str | None = None
    model_year: str | None = None
    collection: str | None = None
    frame_color: str | None = None
    lens_color: str | None = None
    color_code: str | None = None
    material_code: str | None = None
    frame_material: str | None = None
    bridge_size: str | None = None
    branch_size: str | None = None
    uva_filter: str | None = None
    fitting: str | None = None
    release_code: str | None = None
    phase: str | None = None
    order_quantity: float | None = None
    exceeding_quantity: float | None = None
    total_retail_value: float | None = None
    attributes: dict[str, Any] = field(default_factory=dict)
    variants: list[dict[str, Any]] = field(default_factory=list)
    raw_text: str | None = None


def clean_key(value: Any) -> str:
    text = str(value or "").strip().lower().replace("\ufffd", "")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (float, Decimal)) and float(value).is_integer():
        return str(int(value))
    text = " ".join(str(value).replace("\x00", " ").split()).strip()
    return text or None


def as_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    text = re.sub(r"[^0-9,.+-]", "", text)
    if text.endswith("+") and text.count("+") == 1:
        text = text[:-1]
    if not text:
        return None
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    elif "," in text and "." in text:
        text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def canonical_for(header: str) -> str | None:
    key = clean_key(header)
    # Exact non-price names first, then semantic price classification. Supplier
    # files frequently add a currency, MOQ, or internal code to price headings.
    for canonical, aliases in ALIASES.items():
        if key in aliases:
            return canonical
    price_kind = classify_price_header(key)
    if price_kind:
        return price_kind
    return None


def classify_price_header(header: str) -> str | None:
    """Classify supplier price headings without relying on exact wording.

    Only unit prices are normalized. Extended values such as ``Total RRP`` and
    ``Amount (GBP)`` remain available in attributes_json but are deliberately not
    exposed as a unit wholesale/retail price.
    """
    key = clean_key(header)
    words = set(key.split())
    if not key:
        return None
    if "total" in words or key.startswith("amount") or "extended" in words:
        return None

    retail_phrases = (
        "retail", "rrp", "srp", "msrp", "public price", "listing price",
        "selling price", "sale price", "consumer price", "recommended price", "uvp",
    )
    if any(phrase in key for phrase in retail_phrases):
        return "retail_price"

    wholesale_phrases = (
        "wholesale", "whs", "wsp", "net price", "net prices", "your price",
        "trade price", "dealer price", "buying price", "purchase price",
        "supplier price", "unit cost", "cost price", "p tarif", "p tarif",
    )
    if any(phrase in key for phrase in wholesale_phrases):
        return "wholesale_price"

    # A plain Price column, or Price qualified only by currency, is the supplier's
    # offered/unit price in the stock lists handled by this importer.
    non_unit_qualifiers = {"discount", "margin", "markup", "saving", "variance", "difference"}
    if not (words & non_unit_qualifiers) and words & {"price", "prices", "cost", "costs", "tarif", "tariff"}:
        return "wholesale_price"
    return None


def unique_headers(values: Iterable[Any]) -> list[str]:
    result, used = [], {}
    for index, value in enumerate(values, 1):
        base = clean_text(value) or f"column_{index}"
        used[base] = used.get(base, 0) + 1
        result.append(base if used[base] == 1 else f"{base}_{used[base]}")
    return result


def find_header(rows: list[tuple[Any, ...]]) -> int | None:
    best: tuple[float, int] | None = None
    for index, row in enumerate(rows):
        nonempty = [v for v in row if clean_text(v)]
        if len(nonempty) < 2:
            continue
        recognized = sum(canonical_for(clean_text(v) or "") is not None for v in nonempty)
        identity = sum(canonical_for(clean_text(v) or "") in {"name", "model", "sku", "barcode", "brand"} for v in nonempty)
        score = recognized * 10 + identity * 5 + min(len(nonempty), 10) - index * .05
        if identity and (best is None or score > best[0]):
            best = (score, index)
    return best[1] if best else None


def first_value(data: dict[str, Any], canonical: str) -> Any:
    for header, value in data.items():
        if canonical_for(header) == canonical and clean_text(value):
            return value
    return None


def infer_currency(headers: Iterable[str], values: Iterable[Any]) -> str | None:
    text = " ".join([*(str(x) for x in headers), *(str(x) for x in values if x is not None)])
    if "£" in text or "GBP" in text.upper():
        return "GBP"
    if "€" in text or "EUR" in text.upper():
        return "EUR"
    if "$" in text or "USD" in text.upper():
        return "USD"
    return None


def normalized_identifier(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", clean_text(value).upper()) if clean_text(value) else ""


def is_size_header(header: str) -> bool:
    text = str(header).strip().upper().replace(",", ".")
    alpha_sizes = {"XXXS", "XXS", "XS", "S", "M", "L", "XL", "XXL", "XXXL", "3XL", "4XL", "5XL", "6XL", "OS", "ONE SIZE"}
    return text in alpha_sizes or bool(re.fullmatch(r"(?:[A-Z]{1,2})?\d{1,3}(?:\.\d+)?(?:\s*/\s*\d+)?", text))


def parse_combined_model(value: Any) -> dict[str, str]:
    """Split descriptions such as 'Coat, Ref.: X, Colour: navy, Size: 38'."""
    text = clean_text(value) or ""
    result: dict[str, str] = {}
    if not text:
        return result
    name = re.split(r",\s*(?:Ref\.?|Reference)\s*:", text, maxsplit=1, flags=re.I)[0].strip()
    if name and name != text:
        result["name"] = name
    ref = re.search(r"(?:Ref\.?|Reference)\s*:\s*([^,]+)", text, re.I)
    color = re.search(r"(?:Colou?r|Farbe)\s*:\s*([^,]+)", text, re.I)
    size = re.search(r"(?:Size|Gr.{0,3}e)\s*:\s*([^,]+)\s*$", text, re.I)
    if ref: result["model"] = ref.group(1).strip()
    if color: result["color"] = color.group(1).strip()
    if size: result["size"] = size.group(1).strip()
    return result


def infer_gender(text: Any) -> str | None:
    value = f" {clean_key(text)} "
    if any(word in value for word in (" women ", " womens ", " woman ", " wmns ", " damen ")):
        return "Women"
    if any(word in value for word in (" men ", " mens ", " man ", " herren ")):
        return "Men"
    if any(word in value for word in (" kids ", " kid ", " junior ", " jungen ", " kinder ")):
        return "Kids"
    if " unisex " in value:
        return "Unisex"
    return None


def build_auxiliary_maps(workbook: Any) -> dict[str, dict[str, Any]]:
    """Read EAN/HS helper sheets used to enrich, not duplicate, main products."""
    auxiliary: dict[str, dict[str, Any]] = {}
    for sheet in workbook.worksheets:
        if "ean" not in clean_key(sheet.title):
            continue
        preview = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 60), values_only=True))
        header_index = find_header(preview)
        if header_index is None:
            continue
        headers = unique_headers(preview[header_index])
        for row in sheet.iter_rows(min_row=header_index + 2, values_only=True):
            values = list(row[:len(headers)])
            data = {h: json_value(v) for h, v in zip(headers, values) if clean_text(v)}
            sku = first_value(data, "sku")
            if not sku:
                continue
            key = normalized_identifier(sku)
            entry = auxiliary.setdefault(key, {"variants": {}, "attributes": {}})
            for field_name in ("origin", "material", "hs_code"):
                value = first_value(data, field_name)
                if clean_text(value):
                    entry[field_name] = value
            for header, value in data.items():
                if is_size_header(header) and clean_text(value):
                    entry["variants"][str(header)] = clean_text(value)
            entry["attributes"].update(data)
    return auxiliary


def excel_products(path: Path) -> Iterator[Product]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        auxiliary = build_auxiliary_maps(workbook)
        for sheet in workbook.worksheets:
            if clean_key(sheet.title) in SKIP_SHEETS:
                continue
            # EAN sheets either mirror the main list or enrich its size variants.
            if "ean" in clean_key(sheet.title):
                continue
            preview = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 60), values_only=True))
            header_index = find_header(preview)
            if header_index is None:
                continue
            headers = unique_headers(preview[header_index])
            empty_run = 0
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=header_index + 2, values_only=True), header_index + 2
            ):
                values = list(row[: len(headers)])
                if not any(clean_text(v) for v in values):
                    empty_run += 1
                    if empty_run >= 100:
                        break
                    continue
                empty_run = 0
                data = {h: json_value(v) for h, v in zip(headers, values) if clean_text(v)}
                mapped = {key: first_value(data, key) for key in ALIASES}
                if not any(clean_text(mapped[k]) for k in ("name", "model", "sku", "barcode")):
                    continue
                combined = parse_combined_model(mapped["model"])
                # A model is the most reliable name when the supplier has no title.
                name = clean_text(mapped["name"] or combined.get("name") or mapped["model"] or mapped["sku"])
                model = clean_text(combined.get("model") or mapped["model"])
                sku = clean_text(mapped["sku"])
                barcode = clean_text(mapped["barcode"])
                # Excel often converts long identifiers to numeric values.
                for label, value in (("sku", mapped["sku"]), ("barcode", mapped["barcode"])):
                    if isinstance(value, float) and value.is_integer():
                        if label == "sku": sku = str(int(value))
                        else: barcode = str(int(value))
                aux = auxiliary.get(normalized_identifier(sku), {}) if sku else {}
                frame_color = clean_text(mapped["frame_color"])
                description = clean_text(mapped["description"])
                raw_color = clean_text(mapped["color"] or combined.get("color") or frame_color)
                color_code = clean_text(mapped["color_code"])
                if description and raw_color and len(raw_color) <= 6 and re.fullmatch(r"[A-Za-z0-9./-]+", raw_color):
                    color_code = color_code or raw_color
                    color = description
                else:
                    color = raw_color or description
                variants = []
                aux_variants = aux.get("variants", {})
                for header, value in data.items():
                    if not is_size_header(header):
                        continue
                    quantity = as_number(value)
                    if quantity is None or quantity <= 0:
                        continue
                    variants.append({
                        "size": str(header), "quantity": quantity,
                        "barcode": aux_variants.get(str(header)), "attributes": {"source_column": header},
                    })
                if aux.get("attributes"):
                    data["Auxiliary EAN data"] = aux["attributes"]
                source_key = f"{sheet.title}:{row_number}"
                product = Product(
                    source_location=f"sheet={sheet.title};row={row_number}", source_key=source_key,
                    brand=clean_text(mapped["brand"]), name=name, category=clean_text(mapped["category"]),
                    model=model, sku=sku, barcode=barcode, description=description,
                    color=color, size=clean_text(mapped["size"] or combined.get("size")),
                    gender=clean_text(mapped["gender"]) or infer_gender(name), material=clean_text(mapped["material"]),
                    quantity=as_number(mapped["quantity"]), wholesale_price=as_number(mapped["wholesale_price"]),
                    retail_price=as_number(mapped["retail_price"]), currency=infer_currency(headers, values),
                    status=clean_text(mapped["status"]),
                    offer_type=clean_text(mapped["offer_type"]), origin=clean_text(mapped["origin"] or aux.get("origin")),
                    hs_code=clean_text(mapped["hs_code"] or aux.get("hs_code")), grade=clean_text(mapped["grade"]),
                    model_year=clean_text(mapped["model_year"]), collection=clean_text(mapped["collection"]),
                    frame_color=frame_color, lens_color=clean_text(mapped["lens_color"]),
                    color_code=color_code, material_code=clean_text(mapped["material_code"]),
                    frame_material=clean_text(mapped["frame_material"]), bridge_size=clean_text(mapped["bridge_size"]),
                    branch_size=clean_text(mapped["branch_size"]), uva_filter=clean_text(mapped["uva_filter"]),
                    fitting=clean_text(mapped["fitting"]), release_code=clean_text(mapped["release_code"]),
                    phase=clean_text(mapped["phase"]), order_quantity=as_number(mapped["order_quantity"]),
                    exceeding_quantity=as_number(mapped["exceeding_quantity"]),
                    total_retail_value=as_number(mapped["total_retail_value"]),
                    attributes=data, variants=variants,
                )
                if not product.material and aux.get("material"):
                    product.material = clean_text(aux["material"])
                yield product
    finally:
        workbook.close()


def currency_from(text: str) -> str | None:
    if "£" in text: return "GBP"
    if "€" in text: return "EUR"
    if "$" in text: return "USD"
    return None


def parse_offer_page(text: str, page_number: int) -> list[Product]:
    """Parse catalogue entries containing WSP/WHS and RRP price markers."""
    normalized = text.replace("\ufffd", "£")
    # Tempered matching allows line breaks but cannot cross an earlier price
    # marker, keeping each match attached to its own product.
    pattern = re.compile(
        rf"(?P<prefix>(?:(?!(?:WSP|WHS|RRP)\b)[\s\S]){{1,240}}?)\s*"
        rf"(?:WSP|WHS)\s*{MONEY}\s*(?:[-–]\s*)?RRP\s*{MONEY}", re.I
    )
    products = []
    for index, match in enumerate(pattern.finditer(normalized), 1):
        prefix = " ".join(match.group("prefix").split())
        codes = list(PRODUCT_CODE.finditer(prefix))
        if not codes:
            continue
        code_match = codes[-1]
        sku = code_match.group(0).strip("- ")
        before = prefix[: code_match.start()].strip(" -")
        after = prefix[code_match.end():].strip(" -")
        name = after or before[-140:] or sku
        products.append(Product(
            source_location=f"page={page_number}", source_key=f"page:{page_number}:offer:{index}:{sku}",
            name=name, sku=sku, wholesale_price=as_number(match.group(2)),
            retail_price=as_number(match.group(3)), currency=currency_from(match.group(0)) or "GBP",
            raw_text=" ".join(match.group(0).split()), attributes={"parser": "offer_prices"},
        ))
    return products


def parse_lot_page(text: str, page_number: int, unit_price: float | None = None) -> list[Product]:
    """Parse take-all stock lists whose product lines end in Ref/Gamme/Total."""
    products: list[Product] = []
    for i, raw_line in enumerate(text.splitlines()):
        line = " ".join(raw_line.split())
        if not re.search(r"Ref\s*:\s*Gamme\s*:", line, re.I):
            continue
        total = re.search(r"Total\s*(\d+)", line, re.I)
        head = re.split(r"\s*-\s*[^-]*?Ref\s*:", line, maxsplit=1, flags=re.I)[0]
        # Supplier references are fused to the end of the name in these PDFs.
        candidates = list(re.finditer(r"[A-Z]{0,5}\d{4,}[A-Z0-9]*", head, re.I))
        if not candidates:
            continue
        cm = candidates[-1]
        fused = cm.group(0)
        # Prefer the numeric style portion when a preceding name fragment (FG/AG)
        # has become attached to an all-numeric Puma reference.
        numeric_tail = re.search(r"\d{6,}$", fused)
        sku = numeric_tail.group(0) if numeric_tail else fused
        name = head[:cm.start()].strip(" -") or sku
        products.append(Product(
            source_location=f"page={page_number}", source_key=f"page:{page_number}:lot:{i}:{sku}",
            name=name, sku=sku, quantity=as_number(total.group(1)) if total else None,
            wholesale_price=unit_price, currency="EUR" if unit_price is not None else None,
            raw_text=line, attributes={"parser": "take_all_lot"},
        ))
    return products


def parse_destock_page(text: str, page_number: int) -> list[Product]:
    """Parse French stock sheets with P.Tarif, RRP and Total markers."""
    lines = [" ".join(x.split()) for x in text.splitlines() if x.strip()]
    products: list[Product] = []
    for i, line in enumerate(lines):
        if "P.Tarif" not in line or i == 0:
            continue
        detail = lines[i - 1]
        price = re.search(rf"P\.Tarif\s*{MONEY}", line, re.I)
        rrp_line = lines[i + 1] if i + 1 < len(lines) else ""
        rrp = re.search(rf"RRP\s*{MONEY}", rrp_line, re.I)
        total = re.search(r"Total\s*(\d+)", detail, re.I)
        before_ref = re.split(r"\s*-\s*[^-]*?Ref\s*:", detail, maxsplit=1, flags=re.I)[0]
        codes = list(PRODUCT_CODE.finditer(before_ref))
        if not codes or not price:
            continue
        cm = codes[-1]
        sku = cm.group(0)
        name = before_ref[:cm.start()].strip(" -")[-160:] or sku
        products.append(Product(
            source_location=f"page={page_number}", source_key=f"page:{page_number}:destock:{i}:{sku}",
            name=name, sku=sku, quantity=as_number(total.group(1)) if total else None,
            wholesale_price=as_number(price.group(1)), retail_price=as_number(rrp.group(1)) if rrp else None,
            currency=currency_from(line + rrp_line) or "EUR", raw_text=" ".join((detail, line, rrp_line)),
            attributes={"parser": "destock"},
        ))
    return products


def parse_classic_page(text: str, page_number: int) -> list[Product]:
    """Parse footwear catalogue blocks: numeric style, REGULAR/NARROW, price and size."""
    lines = [" ".join(x.split()) for x in text.splitlines() if x.strip()]
    products: list[Product] = []
    for i, line in enumerate(lines):
        if not re.fullmatch(r"\d{6,9}", line) or i + 2 >= len(lines):
            continue
        following = lines[i + 1:i + 4]
        price_line_index = next((j for j, x in enumerate(following) if re.match(r"\d+[.,]\d{2}\s+\d", x)), None)
        if price_line_index is None:
            continue
        price_line = following[price_line_index]
        pm = re.match(r"(\d+[.,]\d{2})\s+(.+)", price_line)
        start = max(0, i - 4)
        name_parts = [x for x in lines[start:i] if not re.search(r"REGULAR|NARROW|Page", x, re.I)]
        name = " ".join(name_parts[-4:]) or line
        variants = {}
        for variant_line in following[:price_line_index]:
            vm = re.match(r"(REGULAR|NARROW):\s*(\d+)", variant_line, re.I)
            if vm: variants[vm.group(1).lower()] = vm.group(2)
        products.append(Product(
            source_location=f"page={page_number}", source_key=f"page:{page_number}:classic:{i}:{line}",
            name=name, model=line, sku=line, size=pm.group(2) if pm else None,
            retail_price=as_number(pm.group(1)) if pm else None, currency="EUR",
            raw_text=" | ".join(lines[i:i + 1 + len(following)]),
            attributes={"parser": "classic_footwear", **variants},
        ))
    return products


def pdf_products(path: Path) -> Iterator[Product]:
    reader = PdfReader(path)
    filename_price = re.search(r"-\s*(\d+(?:[.,]\d+)?)\s*€", path.name)
    lot_price = as_number(filename_price.group(1)) if filename_price else None
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        found = parse_offer_page(text, page_number)
        if not found:
            found = parse_destock_page(text, page_number)
        if not found:
            found = parse_classic_page(text, page_number)
        if not found:
            found = parse_lot_page(text, page_number, lot_price)
        yield from found


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def safe_filename_part(value: Any) -> str:
    text = re.sub(r"[^A-Za-z0-9._-]+", "-", clean_text(value) or "unknown").strip("-.")
    return text[:80] or "unknown"


def resolve_zip_target(base_path: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_path), target))


def relationship_map(archive: zipfile.ZipFile, owner_path: str) -> dict[str, tuple[str, str]]:
    rel_path = posixpath.join(
        posixpath.dirname(owner_path), "_rels", posixpath.basename(owner_path) + ".rels"
    )
    if rel_path not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(rel_path))
    result = {}
    for rel in root:
        rel_id = rel.attrib.get("Id")
        target = rel.attrib.get("Target")
        if rel_id and target and rel.attrib.get("TargetMode") != "External":
            result[rel_id] = (resolve_zip_target(owner_path, target), rel.attrib.get("Type", ""))
    return result


def xlsx_embedded_images(path: Path) -> Iterator[dict[str, Any]]:
    """Yield embedded images with their worksheet cell anchors using OOXML."""
    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_draw = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    ns_art = "http://schemas.openxmlformats.org/drawingml/2006/main"
    with zipfile.ZipFile(path) as archive:
        workbook_path = "xl/workbook.xml"
        workbook = ET.fromstring(archive.read(workbook_path))
        workbook_rels = relationship_map(archive, workbook_path)
        for sheet_node in workbook.findall(f".//{{{ns_main}}}sheet"):
            sheet_name = sheet_node.attrib.get("name", "Sheet")
            rel_id = sheet_node.attrib.get(f"{{{ns_rel}}}id")
            if not rel_id or rel_id not in workbook_rels:
                continue
            sheet_path = workbook_rels[rel_id][0]
            for drawing_path, rel_type in relationship_map(archive, sheet_path).values():
                if not rel_type.endswith("/drawing") or drawing_path not in archive.namelist():
                    continue
                drawing = ET.fromstring(archive.read(drawing_path))
                drawing_rels = relationship_map(archive, drawing_path)
                anchors = list(drawing.findall(f"{{{ns_draw}}}oneCellAnchor")) + list(
                    drawing.findall(f"{{{ns_draw}}}twoCellAnchor")
                )
                for anchor in anchors:
                    row_node = anchor.find(f"{{{ns_draw}}}from/{{{ns_draw}}}row")
                    col_node = anchor.find(f"{{{ns_draw}}}from/{{{ns_draw}}}col")
                    blip = anchor.find(f".//{{{ns_art}}}blip")
                    image_rel_id = blip.attrib.get(f"{{{ns_rel}}}embed") if blip is not None else None
                    if row_node is None or image_rel_id not in drawing_rels:
                        continue
                    image_path = drawing_rels[image_rel_id][0]
                    if image_path not in archive.namelist():
                        continue
                    extent = anchor.find(f"{{{ns_draw}}}ext")
                    yield {
                        "sheet": sheet_name, "row": int(row_node.text or 0) + 1,
                        "column": int(col_node.text or 0) + 1 if col_node is not None else None,
                        "width": int(extent.attrib.get("cx", 0)) / 9525 if extent is not None else None,
                        "height": int(extent.attrib.get("cy", 0)) / 9525 if extent is not None else None,
                        "source_name": posixpath.basename(image_path), "data": archive.read(image_path),
                    }


def stored_media_path(media_root: Path, database_parent: Path, file_path: Path) -> str:
    try:
        return file_path.resolve().relative_to(database_parent.resolve()).as_posix()
    except ValueError:
        return str(file_path.resolve())


def mime_for_extension(extension: str) -> str:
    return {
        ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
        ".gif": "image/gif", ".bmp": "image/bmp", ".tif": "image/tiff",
        ".tiff": "image/tiff", ".webp": "image/webp", ".emf": "image/emf",
        ".wmf": "image/wmf", ".jp2": "image/jp2",
    }.get(extension.lower(), "application/octet-stream")


def product_rows_for_source(connection: sqlite3.Connection, source_id: int) -> dict[str, list[tuple[int, int]]]:
    sheets: dict[str, list[tuple[int, int]]] = {}
    for product_id, location in connection.execute(
        "SELECT id,source_location FROM products WHERE source_file_id=?", (source_id,)
    ):
        match = re.fullmatch(r"sheet=(.*);row=(\d+)", location)
        if match:
            sheets.setdefault(match.group(1), []).append((int(match.group(2)), product_id))
    for rows in sheets.values():
        rows.sort()
    return sheets


def nearest_excel_product(rows: list[tuple[int, int]], anchor_row: int, tolerance: int = 2) -> int | None:
    candidates = [(abs(row - anchor_row), row, product_id) for row, product_id in rows]
    if not candidates:
        return None
    distance, _, product_id = min(candidates)
    return product_id if distance <= tolerance else None


def add_product_image(
    connection: sqlite3.Connection, *, product_id: int, source_id: int, image_path: str,
    digest: str, mime_type: str, source_sheet: str | None = None, source_row: int | None = None,
    source_page: int | None = None, x: float | None = None, y: float | None = None,
    width: float | None = None, height: float | None = None, match_method: str,
) -> None:
    existing_count = connection.execute(
        "SELECT COUNT(*) FROM product_images WHERE product_id=?", (product_id,)
    ).fetchone()[0]
    is_primary = 1 if existing_count == 0 else 0
    connection.execute(
        "INSERT OR IGNORE INTO product_images(product_id,source_file_id,image_path,mime_type,"
        "source_sheet,source_row,source_page,source_x,source_y,width,height,sha256,sort_order,is_primary,match_method) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (product_id, source_id, image_path, mime_type, source_sheet, source_row, source_page,
         x, y, width, height, digest, existing_count, is_primary, match_method),
    )
    if is_primary:
        connection.execute("UPDATE products SET image_path=? WHERE id=?", (image_path, product_id))


def sync_excel_images(
    connection: sqlite3.Connection, source_id: int, path: Path, media_root: Path,
    database_parent: Path,
) -> tuple[int, int]:
    rows_by_sheet = product_rows_for_source(connection, source_id)
    matched = unmatched = 0
    output_dir = media_root / f"source-{source_id:04d}-{safe_filename_part(path.stem)}"
    for image_number, item in enumerate(xlsx_embedded_images(path), 1):
        product_id = nearest_excel_product(rows_by_sheet.get(item["sheet"], []), item["row"])
        if product_id is None:
            unmatched += 1
            continue
        data = item["data"]
        digest = hashlib.sha256(data).hexdigest()
        extension = Path(item["source_name"]).suffix.lower() or ".bin"
        filename = (
            f"{safe_filename_part(item['sheet'])}-row-{item['row']}-col-{item['column'] or 0}-"
            f"{digest[:12]}{extension}"
        )
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / filename
        if not output_path.exists():
            output_path.write_bytes(data)
        stored = stored_media_path(media_root, database_parent, output_path)
        add_product_image(
            connection, product_id=product_id, source_id=source_id, image_path=stored,
            digest=digest, mime_type=mime_for_extension(extension), source_sheet=item["sheet"],
            source_row=item["row"], width=item["width"], height=item["height"],
            match_method="excel_cell_anchor",
        )
        matched += 1
    return matched, unmatched


def pdf_product_positions(page: Any, products: list[sqlite3.Row]) -> list[tuple[sqlite3.Row, Any]]:
    positions = []
    occurrences: dict[str, int] = {}
    for product in products:
        terms = [product["sku"], product["model"], product["name"]]
        rectangles = []
        selected_term = ""
        for term in terms:
            if not term or len(str(term)) < 4:
                continue
            rectangles = page.search_for(str(term))
            if rectangles:
                selected_term = str(term)
                break
        if not rectangles:
            continue
        occurrence = occurrences.get(selected_term, 0)
        rect = rectangles[min(occurrence, len(rectangles) - 1)]
        occurrences[selected_term] = occurrence + 1
        positions.append((product, rect))
    return positions


def sync_pdf_images(
    connection: sqlite3.Connection, source_id: int, path: Path, media_root: Path,
    database_parent: Path,
) -> tuple[int, int]:
    try:
        import pymupdf as fitz  # coordinate-aware PDF image extraction
    except ImportError as exc:
        raise RuntimeError("PyMuPDF is required for PDF image extraction") from exc
    connection.row_factory = sqlite3.Row
    products_by_page: dict[int, list[sqlite3.Row]] = {}
    for product in connection.execute(
        "SELECT id,source_location,source_key,name,model,sku FROM products WHERE source_file_id=? ORDER BY source_key",
        (source_id,),
    ):
        match = re.fullmatch(r"page=(\d+)", product["source_location"])
        if match:
            products_by_page.setdefault(int(match.group(1)), []).append(product)
    matched = unmatched = 0
    output_dir = media_root / f"source-{source_id:04d}-{safe_filename_part(path.stem)}"
    document = fitz.open(str(path))
    try:
        for page_number, products in products_by_page.items():
            if page_number < 1 or page_number > document.page_count:
                continue
            page = document[page_number - 1]
            positions = pdf_product_positions(page, products)
            blocks = [b for b in page.get_text("dict").get("blocks", []) if b.get("type") == 1 and b.get("image")]
            usable = []
            page_area = max(page.rect.width * page.rect.height, 1)
            for block in blocks:
                x0, y0, x1, y1 = block["bbox"]
                width, height = x1 - x0, y1 - y0
                if width < 35 or height < 35 or len(block["image"]) < 1000:
                    continue
                if len(products) > 1 and width * height / page_area > .78:
                    continue
                usable.append(block)
            if not positions:
                unmatched += len(usable)
                continue
            unused = set(range(len(usable)))
            for product, product_rect in positions:
                if not unused:
                    break
                product_center = ((product_rect.x0 + product_rect.x1) / 2, (product_rect.y0 + product_rect.y1) / 2)
                image_number = min(
                    unused,
                    key=lambda index: ((usable[index]["bbox"][0] + usable[index]["bbox"][2]) / 2 - product_center[0]) ** 2
                    + 2 * ((usable[index]["bbox"][1] + usable[index]["bbox"][3]) / 2 - product_center[1]) ** 2,
                )
                block = usable[image_number]
                x0, y0, x1, y1 = block["bbox"]
                image_center = ((x0 + x1) / 2, (y0 + y1) / 2)
                distance = (
                    ((product_rect.x0 + product_rect.x1) / 2 - image_center[0]) ** 2
                    + ((product_rect.y0 + product_rect.y1) / 2 - image_center[1]) ** 2
                ) ** .5
                if distance > ((page.rect.width ** 2 + page.rect.height ** 2) ** .5) * .48:
                    continue
                unused.remove(image_number)
                data = block["image"]
                digest = hashlib.sha256(data).hexdigest()
                extension = "." + str(block.get("ext") or "bin").lower().lstrip(".")
                filename = f"page-{page_number}-image-{image_number + 1}-{digest[:12]}{extension}"
                output_dir.mkdir(parents=True, exist_ok=True)
                output_path = output_dir / filename
                if not output_path.exists():
                    output_path.write_bytes(data)
                stored = stored_media_path(media_root, database_parent, output_path)
                add_product_image(
                    connection, product_id=product["id"], source_id=source_id, image_path=stored,
                    digest=digest, mime_type=mime_for_extension(extension), source_page=page_number,
                    x=x0, y=y0, width=x1 - x0, height=y1 - y0, match_method="pdf_nearest_sku",
                )
                matched += 1
            unmatched += len(unused)
    finally:
        document.close()
    return matched, unmatched


def sync_source_images(
    connection: sqlite3.Connection, source_id: int, path: Path, media_root: Path,
    database_parent: Path,
) -> tuple[int, int]:
    with connection:
        connection.execute("DELETE FROM product_images WHERE source_file_id=?", (source_id,))
        connection.execute("UPDATE products SET image_path=NULL WHERE source_file_id=?", (source_id,))
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            return sync_excel_images(connection, source_id, path, media_root, database_parent)
        if path.suffix.lower() == ".pdf":
            return sync_pdf_images(connection, source_id, path, media_root, database_parent)
    return 0, 0


def infer_brand_from_filename(filename: str) -> str | None:
    """Return a brand only when a source filename is unambiguously single-brand."""
    key = clean_key(filename)
    rules = (
        ("tommy", "TOMMY HILFIGER"), ("thu ", "TOMMY HILFIGER"),
        ("calvin klein", "CALVIN KLEIN"), ("ck ", "CALVIN KLEIN"),
        ("lacoste", "LACOSTE"), ("adidas", "ADIDAS"), ("puma", "PUMA"),
        ("reebok", "REEBOK"), ("asics", "ASICS"), ("timberland", "TIMBERLAND"),
        ("sergio tacchini", "SERGIO TACCHINI"), ("max mara", "MAX MARA"),
        ("helly hansen", "HELLY HANSEN"), ("maserati", "MASERATI"),
    )
    for marker, brand in rules:
        if marker in key:
            return brand
    if "classic women" in key or "classic unisex" in key or "eva styles" in key:
        return "BIRKENSTOCK"
    return None


def insert_product(connection: sqlite3.Connection, source_id: int, p: Product) -> int:
    columns = [
        "source_location", "source_key", "brand", "name", "category", "model", "sku", "barcode",
        "description", "color", "size", "gender", "material", "quantity", "wholesale_price",
        "retail_price", "currency", "status", "image_path", "offer_type", "origin", "hs_code",
        "grade", "model_year", "collection", "frame_color", "lens_color", "color_code",
        "material_code", "frame_material", "bridge_size", "branch_size", "uva_filter", "fitting",
        "release_code", "phase", "order_quantity", "exceeding_quantity", "total_retail_value",
        "attributes_json", "raw_text",
    ]
    values = [getattr(p, c) for c in columns[:-2]] + [
        json.dumps(p.attributes, ensure_ascii=False, default=str, separators=(",", ":")), p.raw_text
    ]
    marks = ",".join("?" for _ in range(len(columns) + 1))
    updates = ",".join(
        f"{column}=excluded.{column}" for column in columns if column not in {"source_key"}
    )
    connection.execute(
        f"INSERT INTO products(source_file_id,{','.join(columns)}) VALUES({marks}) "
        f"ON CONFLICT(source_file_id,source_key) DO UPDATE SET {updates},updated_at=CURRENT_TIMESTAMP",
        [source_id, *values],
    )
    product_id = connection.execute(
        "SELECT id FROM products WHERE source_file_id=? AND source_key=?", (source_id, p.source_key)
    ).fetchone()[0]
    connection.execute("DELETE FROM product_variants WHERE product_id=?", (product_id,))
    for variant in p.variants:
        connection.execute(
            "INSERT INTO product_variants(product_id,size,quantity,barcode,sku,attributes_json) "
            "VALUES(?,?,?,?,?,?) ON CONFLICT(product_id,size) DO UPDATE SET "
            "quantity=excluded.quantity,barcode=excluded.barcode,sku=excluded.sku,attributes_json=excluded.attributes_json",
            (product_id, clean_text(variant.get("size")), as_number(variant.get("quantity")),
             clean_text(variant.get("barcode")), clean_text(variant.get("sku")),
             json.dumps(variant.get("attributes", {}), ensure_ascii=False, default=str, separators=(",", ":"))),
        )
    return product_id


def import_file(
    connection: sqlite3.Connection, path: Path, force: bool = False,
    default_currency: str | None = None,
) -> tuple[str, int, str | None]:
    resolved = str(path.resolve())
    digest = sha256(path)
    existing = connection.execute("SELECT id, sha256, product_count FROM source_files WHERE path=?", (resolved,)).fetchone()
    if existing and existing[1] == digest and not force:
        return "skipped", existing[2], None
    parser = excel_products if path.suffix.lower() in {".xlsx", ".xlsm"} else pdf_products
    try:
        products = parser(path)
        with connection:
            if existing:
                source_id = existing[0]
                connection.execute(
                    "UPDATE source_files SET filename=?,file_type=?,sha256=?,size_bytes=?,modified_at=?,imported_at=CURRENT_TIMESTAMP,warning=NULL WHERE id=?",
                    (path.name, path.suffix.lower()[1:], digest, path.stat().st_size,
                     datetime.fromtimestamp(path.stat().st_mtime).isoformat(), source_id),
                )
            else:
                cursor = connection.execute(
                    "INSERT INTO source_files(path,filename,file_type,sha256,size_bytes,modified_at) VALUES(?,?,?,?,?,?)",
                    (resolved, path.name, path.suffix.lower()[1:], digest, path.stat().st_size,
                     datetime.fromtimestamp(path.stat().st_mtime).isoformat()),
                )
                source_id = cursor.lastrowid
            connection.execute("CREATE TEMP TABLE IF NOT EXISTS seen_source_keys(source_key TEXT PRIMARY KEY)")
            connection.execute("DELETE FROM seen_source_keys")
            count = 0
            for product in products:
                if not product.brand:
                    product.brand = infer_brand_from_filename(path.name)
                if not product.currency:
                    product.currency = default_currency
                insert_product(connection, source_id, product)
                connection.execute(
                    "INSERT OR IGNORE INTO seen_source_keys(source_key) VALUES(?)", (product.source_key,)
                )
                count += 1
            # Remove only rows that disappeared from the changed supplier file.
            # Matching rows are updated in place above, preserving their product IDs.
            connection.execute(
                "DELETE FROM products WHERE source_file_id=? "
                "AND source_key NOT IN (SELECT source_key FROM seen_source_keys)", (source_id,)
            )
            warning = None if count else "No recognizable product rows were found"
            connection.execute("UPDATE source_files SET product_count=?,warning=? WHERE id=?", (count, warning, source_id))
        return "imported", count, warning
    except Exception as exc:
        return "error", 0, f"{type(exc).__name__}: {exc}"


def candidate_files(folder: Path, recursive: bool) -> list[Path]:
    iterator = folder.rglob("*") if recursive else folder.glob("*")
    return sorted(
        (p for p in iterator if p.is_file() and p.suffix.lower() in {".xlsx", ".xlsm", ".pdf"} and not p.name.startswith("~$")),
        key=lambda p: p.name.lower(),
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Ingest product Excel files and PDF catalogues into SQLite")
    parser.add_argument("folder", nargs="?", default=".", type=Path, help="folder containing supplier files")
    parser.add_argument("--database", "-d", default="products.db", type=Path, help="SQLite database path")
    parser.add_argument("--recursive", "-r", action="store_true", help="include subfolders")
    parser.add_argument("--force", action="store_true", help="re-import unchanged files")
    parser.add_argument("--default-currency", default="EUR", help="currency for unlabelled prices (default: EUR; use an empty value to disable)")
    parser.add_argument("--media-dir", type=Path, help="image output folder (default: a media folder beside the database)")
    parser.add_argument("--skip-images", action="store_true", help="do not extract or associate embedded images")
    parser.add_argument("--images-only", action="store_true", help="refresh image associations without re-importing product data")
    parser.add_argument("--limit", type=int, help="only process the first N files (useful for testing)")
    args = parser.parse_args(argv)
    folder = args.folder.resolve()
    files = candidate_files(folder, args.recursive)
    if args.limit is not None:
        files = files[: args.limit]
    args.database.parent.mkdir(parents=True, exist_ok=True)
    media_root = (args.media_dir or (args.database.parent / "media")).resolve()
    connection = sqlite3.connect(args.database)
    connection.executescript(SCHEMA)
    ensure_schema(connection)
    errors = 0
    try:
        for index, path in enumerate(files, 1):
            if args.images_only:
                existing = connection.execute("SELECT id,product_count FROM source_files WHERE path=?", (str(path.resolve()),)).fetchone()
                if existing:
                    status, count, warning = "skipped", existing[1], None
                else:
                    status, count, warning = "error", 0, "source has not been imported"
            else:
                status, count, warning = import_file(connection, path, args.force, args.default_currency or None)
            suffix = f" - {warning}" if warning else ""
            image_suffix = ""
            if not args.skip_images and status != "error" and (status == "imported" or args.images_only):
                source_id = connection.execute("SELECT id FROM source_files WHERE path=?", (str(path.resolve()),)).fetchone()[0]
                try:
                    matched, unmatched = sync_source_images(
                        connection, source_id, path, media_root, args.database.resolve().parent
                    )
                    image_suffix = f"; {matched} images matched, {unmatched} unmatched"
                except Exception as exc:
                    image_suffix = f"; image error: {type(exc).__name__}: {exc}"
                    errors += 1
            print(f"[{index}/{len(files)}] {status:8} {count:6} products  {path.name}{suffix}{image_suffix}", flush=True)
            errors += status == "error"
        total = connection.execute("SELECT COUNT(*) FROM products").fetchone()[0]
        print(f"Database: {args.database.resolve()} ({total} product rows)")
    finally:
        connection.close()
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
